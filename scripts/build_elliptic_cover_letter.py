"""Build the Elliptic cover letter as a .docx matching the resume's typography."""
from __future__ import annotations

import sys
from pathlib import Path

import yaml
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from jobsearch.tailor.writer import _strip_ai_tells  # noqa: E402

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
BLUE = RGBColor(0x2E, 0x75, 0xB6)
GRAY = RGBColor(0x59, 0x59, 0x59)

FONT = "Cambria"
NAME_PT = 14
TAGLINE_PT = 9.5
CONTACT_PT = 9.5
BODY_PT = 10


def _run(p, text, *, bold=False, italic=False, size=BODY_PT, color=None):
    text = _strip_ai_tells(text)
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    if color is not None:
        r.font.color.rgb = color


def _hyperlink(p, url, text, *, size=CONTACT_PT, bold=False):
    text = _strip_ai_tells(text)
    part = p.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rFonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rFonts.set(qn(attr), FONT)
    rPr.append(rFonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rPr.append(sz)
    if bold:
        rPr.append(OxmlElement("w:b"))
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2E75B6")
    rPr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    new_run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    new_run.append(t)
    hyperlink.append(new_run)
    p._p.append(hyperlink)


def _spacing(p, *, before=None, after=None, line=None):
    pf = p.paragraph_format
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing = line


PARAGRAPHS = [
    "At Kuehne+Nagel I work full stack on a regulated travel and expense "
    "platform: React and TypeScript on the front end, Java and Spring on "
    "the back end, MongoDB and Oracle behind that. I pick up tickets across "
    "the stack rather than staying on one side. The work I would point to "
    "is a lazy-loading race that was firing 132 redundant network requests "
    "per page, brought down to 9, and a transaction-flow refactor that cut "
    "latency around 30 percent.",

    "I also run a side project, PetGearHub, on close to your stack: "
    "Next.js and React on Vercel, a NestJS API on AWS ECS, Prisma against "
    "PostgreSQL on RDS. I built it alongside Claude Code. AI coding "
    "assistants are part of how I work day to day; at K+N I have also "
    "integrated MCP servers into our Claude assistants so they reason over "
    "real bug-tracker, IDE, and database context, and I built an OCR and "
    "LLM extraction pipeline for receipts with schema validation and "
    "ground-truth checks before any prompt change ships.",

    "What draws me to Elliptic is the problem domain. Three years on the "
    "Lloyds Banking Group account through TCS taught me what software in "
    "regulated finance has to look like to stand up to an audit, and "
    "tools that help investigators trace fund flows and catch financial "
    "crime are the kind of work I would want my code to be doing.",

    "I will be straight about where I would be growing into the role. My "
    "Node and NestJS work has been the side project rather than years of "
    "production tenure, and I have not used a graph visualisation library "
    "in anger. The way I have picked up the rest of my stack has been by "
    "using it on real work until it ships, and I am comfortable doing "
    "the same here.",

    "Thank you for considering my application. I would welcome the chance "
    "to talk.",
]

POSTSCRIPT = (
    "P.S. I used AI to help format and polish this letter. The work, "
    "opinions, and story are mine; I would rather say so plainly than "
    "pretend otherwise."
)


def _contact_line(p, candidate):
    first = True

    def sep():
        nonlocal first
        if not first:
            _run(p, "  |  ", size=CONTACT_PT, color=GRAY)
        first = False

    if candidate["location"]:
        sep(); _run(p, candidate["location"], size=CONTACT_PT, color=GRAY)
    if candidate["phone"]:
        sep(); _run(p, candidate["phone"], size=CONTACT_PT, color=GRAY)
    if candidate["email"]:
        sep(); _hyperlink(p, f"mailto:{candidate['email']}", candidate["email"], size=CONTACT_PT)
    if candidate.get("linkedin_url"):
        sep(); _hyperlink(p, candidate["linkedin_url"], candidate.get("linkedin") or "LinkedIn", size=CONTACT_PT)
    if candidate.get("github_url"):
        sep(); _hyperlink(p, candidate["github_url"], candidate.get("github") or "GitHub", size=CONTACT_PT)


def main():
    profile_path = ROOT / "config" / "profile.yaml"
    candidate = yaml.safe_load(profile_path.read_text(encoding="utf-8"))["candidate"]

    doc = Document()
    section = doc.sections[0]
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)
    section.top_margin = Inches(0.4)
    section.bottom_margin = Inches(0.4)

    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(BODY_PT)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.line_spacing = 1.15

    # Header (matches resume)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _spacing(p, before=0, after=1, line=1.0)
    _run(p, candidate["full_name"], bold=True, size=NAME_PT, color=NAVY)

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _spacing(p, before=0, after=1, line=1.0)
    _run(p, "Full Stack Engineer  -  TypeScript, React, Node.js", bold=True, size=TAGLINE_PT, color=BLUE)

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _spacing(p, before=0, after=10, line=1.0)
    _contact_line(p, candidate)

    # Date
    p = doc.add_paragraph()
    _spacing(p, before=0, after=8, line=1.15)
    _run(p, "10 June 2026", size=BODY_PT)

    # Salutation
    p = doc.add_paragraph()
    _spacing(p, before=0, after=8, line=1.15)
    _run(p, "Dear Hiring Team,", size=BODY_PT)

    # Body paragraphs
    for para in PARAGRAPHS:
        p = doc.add_paragraph()
        _spacing(p, before=0, after=8, line=1.15)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _run(p, para, size=BODY_PT)

    # Sign-off
    p = doc.add_paragraph()
    _spacing(p, before=0, after=2, line=1.15)
    _run(p, "Kind regards,", size=BODY_PT)

    p = doc.add_paragraph()
    _spacing(p, before=0, after=10, line=1.15)
    _run(p, candidate["full_name"], bold=True, size=BODY_PT, color=NAVY)

    # P.S.
    p = doc.add_paragraph()
    _spacing(p, before=0, after=0, line=1.15)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _run(p, POSTSCRIPT, italic=True, size=BODY_PT - 0.5, color=GRAY)

    out_path = ROOT / "output" / "resumes" / "Sai_Jayanthi_Elliptic_Full_Stack_Cover_Letter.docx"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    print(f"Wrote: {out_path}")


if __name__ == "__main__":
    main()
