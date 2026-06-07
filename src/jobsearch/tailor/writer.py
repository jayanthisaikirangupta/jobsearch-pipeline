"""Render the tailored JSON into a one-page OR two-page .docx.

Two layouts, both reverse-engineered from known-good references:

  ONE-PAGE — based on Sai_Kiran_Avaloq_Angular_Resume.docx (1 page, 56 lines)
    - Cambria 10pt body, 0.5"/0.4" margins, single line spacing
    - For tightly-scoped IC roles, mid-junior seniority

  TWO-PAGE — based on Sai_Kiran_AI_Engineer_Resume.docx (2 pages, 90 lines)
    - Calibri 11pt body, 0.625"/0.556" margins, line_spacing 1.1 on bullets
    - For senior/lead/breadth roles where the full skill story matters

The tailor JSON's ``format_choice`` ("one_page" or "two_page") drives which
layout is used. If ``format_choice`` is missing, default to one_page.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# --- shared palette ---------------------------------------------------------
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
BLUE = RGBColor(0x2E, 0x75, 0xB6)
GRAY = RGBColor(0x59, 0x59, 0x59)


# --- low-level helpers (font-name passed in so both layouts can use these) -

def _spacing(p, before: float | None = None, after: float | None = None,
             line: float | None = None) -> None:
    pf = p.paragraph_format
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing = line


def _add_right_tab(p, position_inches: float) -> None:
    p.paragraph_format.tab_stops.add_tab_stop(
        Inches(position_inches), WD_TAB_ALIGNMENT.RIGHT
    )


def _strip_ai_tells(text: str) -> str:
    """Belt-and-suspenders cleanup: even if Claude slips, the rendered .docx
    never contains long-dash AI tells. Em/en-dashes between words become
    a comma+space, anywhere else they collapse to a single hyphen so we
    don't break compound forms unintentionally. Smart quotes also get
    flattened to ASCII for ATS safety.
    """
    if not text:
        return ""
    # Em/en-dash with surrounding spaces (the typical AI parenthetical) -> comma
    out = text
    for ch in ("—", "–"):  # em-dash, en-dash
        out = out.replace(f" {ch} ", ", ")
        out = out.replace(f"{ch} ", ", ")
        out = out.replace(f" {ch}", ",")
        out = out.replace(ch, "-")  # any remaining occurrence -> hyphen
    # ASCII-friendly quotes (ATS-safe)
    out = out.replace("‘", "'").replace("’", "'")
    out = out.replace("“", '"').replace("”", '"')
    return out


def _run(p, text: str, *, font: str, bold: bool = False, italic: bool = False,
         size: float, color: RGBColor | None = None) -> None:
    text = _strip_ai_tells(text)
    r = p.add_run(text)
    r.font.name = font
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    if color is not None:
        r.font.color.rgb = color


def _add_hyperlink(p, url: str, text: str, *, font: str, size: float,
                   bold: bool = False) -> None:
    """Real clickable hyperlink, blue underlined."""
    text = _strip_ai_tells(text)
    part = p.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")

    rFonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rFonts.set(qn(attr), font)
    rPr.append(rFonts)

    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rPr.append(sz)
    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), str(int(size * 2)))
    rPr.append(szCs)

    if bold:
        rPr.append(OxmlElement("w:b"))

    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2E75B6")
    rPr.append(color)

    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)

    new_run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    new_run.append(t)
    hyperlink.append(new_run)
    p._p.append(hyperlink)


def _truncate(text: str, max_chars: int) -> str:
    if not text or len(text) <= max_chars:
        return text or ""
    cut = text[: max_chars - 1]
    sp = cut.rfind(" ")
    if sp > max_chars * 0.6:
        cut = cut[:sp]
    return cut.rstrip(" ,;:.-") + "."


def _setup_doc(font: str, body_pt: float, margins: tuple[float, float, float, float],
               line_spacing: float | None) -> Document:
    doc = Document()
    section = doc.sections[0]
    l, r, t, b = margins
    section.left_margin = Inches(l)
    section.right_margin = Inches(r)
    section.top_margin = Inches(t)
    section.bottom_margin = Inches(b)

    style = doc.styles["Normal"]
    style.font.name = font
    style.font.size = Pt(body_pt)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.line_spacing = line_spacing
    return doc


def _contact_line(p, candidate, *, font: str, size: float) -> None:
    """Center-aligned contact line with real clickable hyperlinks."""
    first = True

    def sep() -> None:
        nonlocal first
        if not first:
            _run(p, "  |  ", font=font, size=size, color=GRAY)
        first = False

    if candidate.location:
        sep(); _run(p, candidate.location, font=font, size=size, color=GRAY)
    if candidate.phone:
        sep(); _run(p, candidate.phone, font=font, size=size, color=GRAY)
    if candidate.email:
        sep(); _add_hyperlink(p, f"mailto:{candidate.email}", candidate.email,
                              font=font, size=size)
    linkedin_url = getattr(candidate, "linkedin_url", "")
    linkedin_label = candidate.linkedin or "LinkedIn"
    if linkedin_url:
        sep(); _add_hyperlink(p, linkedin_url, linkedin_label, font=font, size=size)
    elif candidate.linkedin:
        sep(); _run(p, candidate.linkedin, font=font, size=size, color=GRAY)
    github_url = getattr(candidate, "github_url", "")
    github_label = candidate.github or "GitHub"
    if github_url:
        sep(); _add_hyperlink(p, github_url, github_label, font=font, size=size)
    elif candidate.github:
        sep(); _run(p, candidate.github, font=font, size=size, color=GRAY)


# ===========================================================================
# ONE-PAGE LAYOUT — Avaloq style
# ===========================================================================

ONE_PAGE = {
    "font": "Cambria",
    "name_pt": 14, "tagline_pt": 9.5, "contact_pt": 9.5,
    "header_pt": 10.5, "body_pt": 10, "bullet_pt": 10,
    "margins": (0.5, 0.5, 0.4, 0.4),
    "usable_width": 7.5,
    "line_spacing": None,  # Word default = single
    "max_profile_chars": 480,
    "max_core_skill_rows": 6,
    "max_bullets_per_role": 4,
    # Claude is asked to write <=160 chars; this 200 cap is the safety net for
    # the rare case it overshoots. Keep it loose so we never cut mid-sentence.
    "max_bullet_chars": 200,
    "max_roles": 3,
    "max_projects": 2,
    "max_education": 2,
}


def _render_one_page(tailored: dict[str, Any], candidate, out_path: Path) -> Path:
    cfg = ONE_PAGE
    font = cfg["font"]
    body = cfg["body_pt"]
    doc = _setup_doc(font, body, cfg["margins"], cfg["line_spacing"])

    # Name
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _spacing(p, before=0, after=1)
    _run(p, candidate.full_name, font=font, bold=True, size=cfg["name_pt"], color=NAVY)

    # Tagline
    # JD-adaptive headline takes priority over static profile tagline.
    # Claude's tailored headline is computed per-JD; the candidate's static
    # tagline is the fallback only when an old tailor JSON pre-dates the
    # headline field.
    if tagline := tailored.get("headline") or getattr(candidate, "tagline", ""):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _spacing(p, before=0, after=1)
        _run(p, tagline, font=font, bold=True, size=cfg["tagline_pt"], color=BLUE)

    # Contact
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _spacing(p, before=0, after=3)
    _contact_line(p, candidate, font=font, size=cfg["contact_pt"])

    # Profile
    if profile := tailored.get("profile") or tailored.get("headline"):
        _section_header(doc, "Personal Profile", font=font, size=cfg["header_pt"])
        p = doc.add_paragraph(); _spacing(p, before=0, after=0.5)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _run(p, _truncate(profile, cfg["max_profile_chars"]), font=font, size=body)

    # Experience
    if experience := tailored.get("experience"):
        _section_header(doc, "Professional Experience", font=font, size=cfg["header_pt"])
        for role in experience[:cfg["max_roles"]]:
            p = doc.add_paragraph(); _spacing(p, before=3, after=0.5)
            _add_right_tab(p, cfg["usable_width"])
            title = (role.get("title", "") or "").upper()
            company = role.get("company", "")
            header_text = " - ".join(b for b in [title, company] if b)
            _run(p, header_text, font=font, bold=True, size=body, color=NAVY)
            if dates := role.get("dates", ""):
                _run(p, f"\t{dates}", font=font, size=body, color=GRAY)

            sub_parts = [v for k in ("subtitle", "location") if (v := role.get(k))]
            if sub_parts:
                p = doc.add_paragraph(); _spacing(p, before=0, after=0.5)
                _run(p, "  ·  ".join(sub_parts), font=font, italic=True,
                     size=body - 1, color=GRAY)

            for b in role.get("bullets", [])[:cfg["max_bullets_per_role"]]:
                _bullet(doc, _truncate(b, cfg["max_bullet_chars"]),
                        font=font, size=cfg["bullet_pt"])

    # Skills
    skills = tailored.get("core_skills") or _flatten_skills_grouped(tailored.get("skills_grouped"))
    if skills:
        _section_header(doc, "Skills", font=font, size=cfg["header_pt"])
        for row in skills[:cfg["max_core_skill_rows"]]:
            p = doc.add_paragraph(); _spacing(p, before=0, after=0.5)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _run(p, f"{row['category']}: ", font=font, bold=True, size=body, color=NAVY)
            _run(p, row["items"], font=font, size=body)

    # Projects
    if projects := tailored.get("projects"):
        _section_header(doc, "AI Projects & Practical Work", font=font, size=cfg["header_pt"])
        for proj in projects[:cfg["max_projects"]]:
            p = doc.add_paragraph(); _spacing(p, before=1, after=0)
            _run(p, proj.get("name", ""), font=font, bold=True, size=body, color=NAVY)
            if stack := proj.get("tech_stack"):
                _run(p, f"  ·  {stack}", font=font, italic=True, size=body - 1, color=GRAY)
            if summary := proj.get("summary"):
                _bullet(doc, _truncate(summary, cfg["max_bullet_chars"]),
                        font=font, size=cfg["bullet_pt"])

    # Education
    if education := tailored.get("education"):
        _section_header(doc, "Education", font=font, size=cfg["header_pt"])
        for ed in education[:cfg["max_education"]]:
            p = doc.add_paragraph(); _spacing(p, before=1, after=0.5)
            _add_right_tab(p, cfg["usable_width"])
            _run(p, ed.get("degree", ""), font=font, bold=True, size=body, color=NAVY)
            if inst := ed.get("institution"):
                _run(p, f"  ·  {inst}", font=font, size=body)
            if dates := ed.get("dates"):
                _run(p, f"\t{dates}", font=font, size=body, color=GRAY)

    # Certifications
    if certs := tailored.get("certifications"):
        p = doc.add_paragraph(); _spacing(p, before=2, after=0.5)
        _run(p, "Certifications: ", font=font, bold=True, size=body, color=NAVY)
        _run(p, certs, font=font, size=body)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path


# ===========================================================================
# TWO-PAGE LAYOUT — AI Engineer style
# ===========================================================================

TWO_PAGE = {
    "font": "Calibri",
    "name_pt": 14, "tagline_pt": 11, "contact_pt": 11,
    "header_pt": 12, "body_pt": 11, "bullet_pt": 11,
    "margins": (0.625, 0.625, 0.556, 0.556),
    "usable_width": 7.25,
    "line_spacing": 1.1,
    "max_profile_chars": 700,
    "max_core_skill_rows": 8,
    "max_bullets_per_role": 7,
    # Same safety-net policy as one-page: prompt asks for <=200, this is the
    # cap for the rare overflow case.
    "max_bullet_chars": 240,
    "max_roles": 3,
    "max_projects": 3,
    "max_education": 2,
}


def _render_two_page(tailored: dict[str, Any], candidate, out_path: Path) -> Path:
    cfg = TWO_PAGE
    font = cfg["font"]
    body = cfg["body_pt"]
    doc = _setup_doc(font, body, cfg["margins"], cfg["line_spacing"])

    # Name (centered, 14pt bold, sa=3pt — exact from reference)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _spacing(p, before=0, after=3, line=1.0)
    _run(p, candidate.full_name.upper(), font=font, bold=True,
         size=cfg["name_pt"], color=NAVY)

    # Tagline (centered, 11pt bold blue, sa=3pt)
    # JD-adaptive headline takes priority over static profile tagline.
    # Claude's tailored headline is computed per-JD; the candidate's static
    # tagline is the fallback only when an old tailor JSON pre-dates the
    # headline field.
    if tagline := tailored.get("headline") or getattr(candidate, "tagline", ""):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _spacing(p, before=0, after=3, line=1.0)
        _run(p, tagline, font=font, bold=True, size=cfg["tagline_pt"], color=BLUE)

    # Contact (centered, 11pt gray, sa=7pt — gives breathing room from PROFILE)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _spacing(p, before=0, after=7, line=1.0)
    _contact_line(p, candidate, font=font, size=cfg["contact_pt"])

    # Profile
    if profile := tailored.get("profile") or tailored.get("headline"):
        _section_header(doc, "Profile", font=font, size=cfg["header_pt"], style="two_page")
        p = doc.add_paragraph(); _spacing(p, before=0, after=0, line=1.15)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _run(p, _truncate(profile, cfg["max_profile_chars"]), font=font, size=body)

    # Core skills (8 rows allowed when extended_skills supplied)
    skills = tailored.get("core_skills") or _flatten_skills_grouped(tailored.get("skills_grouped")) or []
    extended = tailored.get("extended_skills") or []
    all_skills = list(skills) + list(extended)
    if all_skills:
        _section_header(doc, "Core Skills", font=font, size=cfg["header_pt"], style="two_page")
        for row in all_skills[:cfg["max_core_skill_rows"]]:
            p = doc.add_paragraph(); _spacing(p, before=0, after=3, line=1.1)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _run(p, f"{row['category']}: ", font=font, bold=True, size=body, color=NAVY)
            _run(p, row["items"], font=font, size=body)

    # Experience — 6-7 bullets per role using extra_bullets
    if experience := tailored.get("experience"):
        _section_header(doc, "Professional Experience", font=font, size=cfg["header_pt"], style="two_page")
        for role in experience[:cfg["max_roles"]]:
            # Line 1: Company \t Dates (matches AI Engineer reference)
            p = doc.add_paragraph(); _spacing(p, before=6, after=1.5, line=1.0)
            _add_right_tab(p, cfg["usable_width"])
            _run(p, role.get("company", ""), font=font, bold=True, size=body, color=NAVY)
            if dates := role.get("dates", ""):
                _run(p, f"\t{dates}", font=font, bold=True, size=body, color=GRAY)

            # Line 2: Role · subtitle \t Location
            p = doc.add_paragraph(); _spacing(p, before=0, after=3, line=1.0)
            _add_right_tab(p, cfg["usable_width"])
            title_bits = [role.get("title", "")]
            if subtitle := role.get("subtitle"):
                title_bits.append(subtitle)
            _run(p, "  ·  ".join(t for t in title_bits if t),
                 font=font, italic=True, size=body, color=GRAY)
            if loc := role.get("location"):
                _run(p, f"\t{loc}", font=font, italic=True, size=body, color=GRAY)

            all_bullets = list(role.get("bullets", []))
            all_bullets.extend(role.get("extra_bullets", []) or [])
            for b in all_bullets[:cfg["max_bullets_per_role"]]:
                _bullet(doc, _truncate(b, cfg["max_bullet_chars"]),
                        font=font, size=cfg["bullet_pt"], line=1.1)

    # Projects — 3rd via third_project
    projects = list(tailored.get("projects") or [])
    if third := tailored.get("third_project"):
        projects.append(third)
    if projects:
        _section_header(doc, "AI Projects & Practical Work",
                        font=font, size=cfg["header_pt"], style="two_page")
        for proj in projects[:cfg["max_projects"]]:
            p = doc.add_paragraph(); _spacing(p, before=3, after=0, line=1.0)
            _run(p, proj.get("name", ""), font=font, bold=True, size=body, color=NAVY)
            if stack := proj.get("tech_stack"):
                _run(p, f"  ·  {stack}", font=font, italic=True, size=body, color=GRAY)
            if summary := proj.get("summary"):
                _bullet(doc, _truncate(summary, cfg["max_bullet_chars"]),
                        font=font, size=cfg["bullet_pt"], line=1.1)

    # Education with modules
    if education := tailored.get("education"):
        _section_header(doc, "Education", font=font, size=cfg["header_pt"], style="two_page")
        for ed in education[:cfg["max_education"]]:
            p = doc.add_paragraph(); _spacing(p, before=3, after=1, line=1.0)
            _add_right_tab(p, cfg["usable_width"])
            _run(p, ed.get("degree", ""), font=font, bold=True, size=body, color=NAVY)
            if inst := ed.get("institution"):
                _run(p, f"  ·  {inst}", font=font, size=body)
            if dates := ed.get("dates"):
                _run(p, f"\t{dates}", font=font, size=body, color=GRAY)
            if modules := ed.get("modules"):
                p = doc.add_paragraph(); _spacing(p, before=0, after=2, line=1.1)
                # Strip a leading 'Modules:' if Claude already wrote it
                modules_text = modules.strip()
                if modules_text.lower().startswith("modules:"):
                    modules_text = modules_text[len("modules:"):].strip()
                _run(p, f"Modules: {modules_text}", font=font, italic=True,
                     size=body - 0.5, color=GRAY)

    # Certifications — detailed line if cert_details provided
    cert_line = tailored.get("cert_details") or tailored.get("certifications")
    if cert_line:
        p = doc.add_paragraph(); _spacing(p, before=4, after=0, line=1.1)
        _run(p, "Certifications: ", font=font, bold=True, size=body, color=NAVY)
        _run(p, cert_line, font=font, size=body)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path


# --- shared section header / bullet (work for both layouts) -----------------

def _add_bottom_rule(p, hex_color: str = "2E75B6") -> None:
    """Add a thin bottom border to a paragraph (used under section headers)."""
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")        # 0.75pt
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), hex_color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def _section_header(doc: Document, title: str, *, font: str, size: float,
                    style: str = "one_page") -> None:
    p = doc.add_paragraph()
    if style == "two_page":
        _spacing(p, before=9, after=5, line=1.0)
    else:
        _spacing(p, before=4, after=1.5, line=1.0)
    _run(p, title.upper(), font=font, bold=True, size=size, color=BLUE)
    _add_bottom_rule(p)


def _bullet(doc: Document, text: str, *, font: str, size: float,
            line: float | None = None) -> None:
    p = doc.add_paragraph(style="List Bullet")
    _spacing(p, before=0, after=0.5 if line is None else 3, line=line)
    p.paragraph_format.left_indent = Inches(0.2)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for r in p.runs:
        r.text = ""
    _run(p, text, font=font, size=size)


def _flatten_skills_grouped(grouped: dict | None) -> list[dict] | None:
    if not grouped:
        return None
    return [{"category": k, "items": ", ".join(v)} for k, v in grouped.items()]


# ===========================================================================
# Public entry point
# ===========================================================================

def render(tailored: dict[str, Any], candidate, out_path: Path) -> Path:
    """Dispatch to the layout chosen by ``tailored['format_choice']``.

    Defaults to one_page when missing or invalid.
    """
    choice = (tailored.get("format_choice") or "one_page").lower().strip()
    if choice == "two_page":
        return _render_two_page(tailored, candidate, out_path)
    return _render_one_page(tailored, candidate, out_path)
